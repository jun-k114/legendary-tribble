"""請求・入金管理表を生成するスクリプト。

使い方:
  python tools/build_workbook.py 出力.xlsx --sample
      動作確認用のサンプル取引先・請求・入金入り
  python tools/build_workbook.py 出力.xlsx --from 現行の収支一覧表.xlsx [--fy-start 2026-10-01]
      現行表の得意先を取引先マスタに登録し、支払条件・毎月請求の有無を入金実績から推定して初期値にする

数式は行番号の位置に「#」を書き、列は列名(キー)から組み立てる。
"""
import argparse
from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ap = argparse.ArgumentParser()
ap.add_argument("out", nargs="?", default="請求入金管理表.xlsx")
ap.add_argument("--sample", action="store_true", help="サンプルデータを入れる")
ap.add_argument("--from", dest="src", help="現行の収支一覧表(取引先マスタの初期値に使う)")
ap.add_argument("--fy-start", default="2026-10-01", help="月別集計の期首(YYYY-MM-01)")
args = ap.parse_args()
FY_START = date.fromisoformat(args.fy_start)

FONT = "Meiryo UI"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True)
F_NOTE = Font(name=FONT, size=9, color="595959")

FILL_HEAD = PatternFill("solid", fgColor="1F4E78")
FILL_HEAD_AUTO = PatternFill("solid", fgColor="595959")
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
FILL_AUTO = PatternFill("solid", fgColor="EDEDED")
FILL_SECTION = PatternFill("solid", fgColor="DDEBF7")

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(vertical="top", wrap_text=True)
UNLOCKED = Protection(locked=False)

FMT_DATE = "yyyy/mm/dd"
FMT_MONTH = "yyyy年m月"
FMT_YEN = "#,##0;[Red]-#,##0"

HDR, FIRST = 4, 5
INV_N, DEP_N, CLI_N = 500, 1000, 200
LAST = {"inv": FIRST + INV_N - 1, "dep": FIRST + DEP_N - 1, "cli": FIRST + CLI_N - 1}

SHEET = {"set": "'設定'", "inv": "'請求台帳'", "dep": "'入金明細'", "cli": "'取引先マスタ'",
         "dash": "'チェック'", "bal": "'取引先別残高'"}
BASE_DATE = f"{SHEET['set']}!$C$5"
NEAR_DAYS = f"{SHEET['set']}!$C$6"
PREFIX = f"{SHEET['set']}!$C$4"
FY_CELL = f"{SHEET['set']}!$C$7"
TARGET_MONTH = f"{SHEET['dash']}!$C$4"

# ---------------------------------------------------------------- 列の定義(キー → 列)
INV_KEYS = ["no", "date", "code", "name", "subject", "amount", "due_auto", "due_manual", "due",
            "method", "sent", "sender", "paid", "last_paid", "adj", "bal", "cancel", "status",
            "over", "missing", "checker", "note"]
DEP_KEYS = ["no", "date", "account", "payer", "amount", "inv", "code", "name", "inv_amount", "state",
            "clerk", "note"]
CLI_KEYS = ["code", "name", "kana", "close", "site", "payday", "monthly", "method", "dest", "rep",
            "note", "dup"]


def letters(keys):
    return {k: get_column_letter(i) for i, k in enumerate(keys, 1)}


I, D, C = letters(INV_KEYS), letters(DEP_KEYS), letters(CLI_KEYS)


def full(sheet, col):
    return f"{SHEET[sheet]}!${col}${FIRST}:${col}${LAST[sheet]}"


INV = {k: full("inv", v) for k, v in I.items()}
DEP = {k: full("dep", v) for k, v in D.items()}
CLI = {k: full("cli", v) for k, v in C.items()}


def cell(cols, key):
    """同じ行のセル参照(行番号は # で後から置換)"""
    return f"{cols[key]}#"


wb = Workbook()


def title(ws, text, note):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    ws["A2"] = note
    ws["A2"].font = F_NOTE
    ws.sheet_view.showGridLines = False


def table(ws, columns, last):
    """columns: (見出し, 幅, 'in'|'auto', 書式, 数式(# = 行) or None, コメント)"""
    for i, (head, width, kind, fmt, formula, comment) in enumerate(columns, start=1):
        col = get_column_letter(i)
        h = ws.cell(row=HDR, column=i, value=head)
        h.font, h.alignment, h.border = F_HEAD, CENTER, BORDER
        h.fill = FILL_HEAD if kind == "in" else FILL_HEAD_AUTO
        if comment:
            h.comment = Comment(comment, "管理表")
        ws.column_dimensions[col].width = width
        for r in range(FIRST, last + 1):
            c = ws.cell(row=r, column=i)
            c.font, c.border = F_BASE, BORDER
            c.fill = FILL_INPUT if kind == "in" else FILL_AUTO
            if fmt:
                c.number_format = fmt
            if formula:
                c.value = formula.replace("#", str(r))
            if kind == "in":
                c.protection = UNLOCKED
    ws.row_dimensions[HDR].height = 32
    ws.freeze_panes = ws.cell(row=FIRST, column=2)
    ws.auto_filter.ref = f"A{HDR}:{get_column_letter(len(columns))}{last}"


def protect(ws):
    # パスワードなし。数式の誤上書きを防ぐ目的(フィルタ・列幅変更は可、並べ替えは不可)
    p = ws.protection
    p.sheet = True
    p.autoFilter = p.formatColumns = p.formatRows = p.formatCells = False


def validation(ws, ref, dv, error):
    dv.error, dv.errorTitle, dv.showErrorMessage = error, "入力エラー", True
    ws.add_data_validation(dv)
    dv.add(ref)


def add_list(ws, ref, source, error="一覧から選択してください"):
    validation(ws, ref, DataValidation(type="list", formula1=source, allow_blank=True), error)


def add_date(ws, ref):
    validation(ws, ref, DataValidation(type="date", operator="between", formula1="DATE(2000,1,1)",
                                       formula2="DATE(2100,12,31)", allow_blank=True),
               "日付を入力してください(例: 2026/9/30)")


def add_number(ws, ref):
    validation(ws, ref, DataValidation(type="decimal", operator="between", formula1="-999999999999",
                                       formula2="999999999999", allow_blank=True), "数値(円)を入力してください")


def highlight(ws, ref, formula, bg, fg, bold=True):
    ws.conditional_formatting.add(ref, FormulaRule(
        formula=[formula], fill=PatternFill("solid", fgColor=bg, bgColor=bg),
        font=Font(name=FONT, bold=bold, color=fg)))


STATUS_COLORS = [
    ("入金済", "E2EFDA", "375623"), ("消込済", "E2EFDA", "375623"),
    ("期日超過", "F8CBAD", "9C0006"), ("一部入金・超過", "F8CBAD", "9C0006"), ("未消込", "F8CBAD", "9C0006"),
    ("未送付", "FCE4D6", "C65911"),
    ("要入力", "FFFF00", "9C0006"), ("番号誤り", "FFFF00", "9C0006"), ("日付誤り", "FFFF00", "9C0006"),
    ("取消請求へ消込", "FFFF00", "9C0006"),
    ("期日間近", "FFE699", "7F6000"), ("一部入金", "FFE699", "7F6000"),
    ("過入金", "E4DFEC", "5B2C6F"), ("取消", "D9D9D9", "7F7F7F"),
]


def status_colors(ws, col, last):
    for text, bg, fg in STATUS_COLORS:
        highlight(ws, f"{col}{FIRST}:{col}{last}", f'${col}{FIRST}="{text}"', bg, fg)


# ---------------------------------------------------------------- はじめに
ws = wb.active
ws.title = "はじめに"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 24
ws.column_dimensions["C"].width = 96
lines = [
    ("title", "請求・入金管理表  ― 使い方"),
    ("note", "請求書の「発行漏れ」「送付漏れ」「入金確認漏れ」「記入漏れ」を仕組みで防ぐための管理表です。金額はすべて税込。"),
    ("", ""),
    ("section", "1. 色の見方"),
    ("legend_in", "黄色のセル", "入力するセル(ここだけ入力します)"),
    ("legend_auto", "灰色のセル", "自動計算(入力不要。シート保護で上書きできないようにしています)"),
    ("", ""),
    ("section", "2. シート構成"),
    ("row", "チェック", "要入力/未送付/期日超過/未消込などの件数。0件でない項目が「要対応」。まずここを開きます。"),
    ("row", "請求台帳", "請求書1枚=1行。請求番号は台帳が自動採番。入金期日・残高・ステータス・入力漏れを自動判定。"),
    ("row", "入金明細", "通帳/ネットバンキングの入金を【全件】転記し、消込先の請求番号を入れます。"),
    ("row", "取引先別残高", "取引先ごとの未回収残高(超過日数別)、毎月請求先の請求漏れ候補。"),
    ("row", "月別集計", "取引先×月の請求額・入金額(収支一覧表の収入欄と同じ形)。収支一覧表への転記・報告用。"),
    ("row", "取引先マスタ", "取引先コード・支払条件(サイト/支払日)・毎月請求の有無を登録。"),
    ("row", "月次チェック記録", "月次締めのチェック結果と承認を記録(誰がいつ確認したかを残す)。"),
    ("row", "設定", "請求番号の接頭辞、基準日、期日間近の日数、月別集計の期首、プルダウンの選択肢。"),
    ("", ""),
    ("section", "3. 基本の流れ"),
    ("row", "① 請求書を作る前", "請求台帳に1行追加(請求日・取引先・件名・請求額・送付方法)→ 採番された請求番号を請求書に記載。"),
    ("row", "② 請求書を送ったら", "送付日・送付者を入力。未入力の間はステータスが「未送付」のまま残ります。"),
    ("row", "③ 入金があったら", "入金明細に入金日・名義・金額を転記し、消込先の請求番号を入力。台帳の入金額・残高が自動更新。"),
    ("row", "④ 消込を確認したら", "残高が0になった請求を別の人が確認し「消込確認者」に記名(未記名だと「要入力」)。"),
    ("row", "⑤ 毎週", "チェックシートで「期日超過」「期日間近」を確認し、担当者から督促・状況確認。"),
    ("row", "⑥ 毎月(締め後)", "チェックシートの全項目が「OK」になることを確認し、月次チェック記録に実施者・承認者を記入。"),
    ("", ""),
    ("section", "4. ステータスの意味(請求台帳)"),
    ("row", "要入力", "必須項目の漏れ、取引先コード誤り、支払条件未登録、未来の送付日、送付者・消込確認者の未記入。「入力漏れ」列に内容を表示。"),
    ("row", "未送付", "台帳登録済みだが送付日が空欄。請求書の出し忘れ・送り忘れの可能性。"),
    ("row", "入金待ち / 期日間近", "送付済みで期日前。期日まで「設定」の日数以内なら期日間近。"),
    ("row", "期日超過", "入金期日を過ぎても未回収残高がある。督促が必要。"),
    ("row", "一部入金 / 一部入金・超過", "入金額が請求額に足りない(振込手数料の差引なら「手数料・値引等」に入力)。"),
    ("row", "過入金", "請求額より多く入金されている(二重振込・消込先誤りの可能性)。"),
    ("row", "入金済", "未回収残高が0。"),
    ("row", "取消", "「取消」列で取消した請求。行は削除せず残します(番号の欠番を作らないため)。"),
    ("", ""),
    ("section", "5. 守るルール"),
    ("row", "請求書1枚=1行", "同じ取引先・同じ月でも請求書が別なら行を分ける(金額を「=A+B」で合算しない)。入金確認は請求書単位。"),
    ("row", "日付は年から入力", "「12/30」だけだと入力した日の年が補われ、年をまたぐと誤る。「2026/12/30」のように年から入力。未来日付はエラー表示。"),
    ("row", "入金は通帳の実額", "入金額は通帳の金額を入力(請求額をそのまま写さない)。1回の入金ごとに1行。"),
    ("row", "入金は全件転記", "請求に対応しない入金も含めて全件入力。「未消込」が残れば消込漏れか不明入金。"),
    ("row", "1入金で複数請求", "入金明細を請求の数だけ行に分け、金額を按分して各請求番号を入力。"),
    ("row", "振込手数料の差引", "先方負担の手数料が差し引かれた場合、請求台帳の「手数料・値引等」に差額を入力。"),
    ("row", "行を削除しない", "間違えた請求は「取消」を選び、備考に理由。正しい請求は新しい行で再発行。"),
    ("row", "並べ替えをしない", "請求番号は行位置から採番されるため並べ替え不可(シート保護で制限)。絞り込みはフィルタで。"),
    ("row", "シート保護", "パスワードなしで保護。列の追加などが必要な場合は[校閲]→[シート保護の解除]。"),
]
r = 1
for item in lines:
    kind = item[0]
    if kind == "title":
        ws.cell(row=r, column=2, value=item[1]).font = F_TITLE
    elif kind == "note":
        ws.cell(row=r, column=2, value=item[1]).font = F_NOTE
    elif kind == "section":
        ws.cell(row=r, column=2, value=item[1]).font = F_BOLD
        for col in (2, 3):
            ws.cell(row=r, column=col).fill = FILL_SECTION
    elif kind:
        a = ws.cell(row=r, column=2, value=item[1])
        b = ws.cell(row=r, column=3, value=item[2])
        a.font, b.font = F_BOLD, F_BASE
        a.alignment = b.alignment = WRAP
        a.border = b.border = BORDER
        if kind == "legend_in":
            a.fill = FILL_INPUT
        elif kind == "legend_auto":
            a.fill = FILL_AUTO
    r += 1

# ---------------------------------------------------------------- 設定
st = wb.create_sheet("設定")
title(st, "設定", "黄色のセルを必要に応じて変更してください。プルダウンの選択肢は各リストの空欄に追記できます。")
for col, w in zip("ABCD", (3, 24, 16, 52)):
    st.column_dimensions[col].width = w
settings = [
    (4, "請求番号の接頭辞", "INV-", None, "請求番号 = 接頭辞 + 5桁の連番(例: INV-00001)"),
    (5, "基準日", "=TODAY()", FMT_DATE, "通常は =TODAY() のまま。過去時点で確認したい場合のみ日付を直接入力"),
    (6, "期日間近とする日数", 7, '0"日"', "入金期日まで何日以内を「期日間近」とするか"),
    (7, "月別集計の期首", FY_START, FMT_DATE, "月別集計シートの1か月目(例: 2026/10/1 = 3期)"),
]
for row, label, val, fmt, note in settings:
    st.cell(row=row, column=2, value=label).font = F_BOLD
    c = st.cell(row=row, column=3, value=val)
    c.font, c.fill, c.border, c.protection = F_BASE, FILL_INPUT, BORDER, UNLOCKED
    if fmt:
        c.number_format = fmt
    st.cell(row=row, column=4, value=note).font = F_NOTE

LIST_FIRST, LIST_LAST = 5, 24
lists = {
    "F": ("送付方法", ["メール(PDF)", "郵送", "請求書発行システム", "先方システムへ登録", "手渡し", "その他"]),
    "G": ("入金口座", ["平塚信用金庫", "(口座名を入力)"]),
    "H": ("取消", ["取消"]),
    "I": ("有無", ["有", "無"]),
    "J": ("実施", ["済", "該当なし"]),
}
for col, (head, values) in lists.items():
    st.column_dimensions[col].width = 20
    h = st[f"{col}4"]
    h.value, h.font, h.fill, h.alignment = head, F_HEAD, FILL_HEAD, CENTER
    for i in range(LIST_FIRST, LIST_LAST + 1):
        c = st[f"{col}{i}"]
        c.font, c.fill, c.border, c.protection = F_BASE, FILL_INPUT, BORDER, UNLOCKED
    for i, v in enumerate(values):
        st[f"{col}{LIST_FIRST + i}"] = v
st["F2"] = "▼ プルダウンの選択肢(空欄に追記すると選択肢に増えます)"
st["F2"].font = F_NOTE
protect(st)


def list_src(col):
    return f"{SHEET['set']}!${col}${LIST_FIRST}:${col}${LIST_LAST}"


# ---------------------------------------------------------------- 取引先マスタ
cm = wb.create_sheet("取引先マスタ")
title(cm, "取引先マスタ",
      "取引先ごとの支払条件を登録します。請求台帳の入金期日は「支払サイト」「支払日」から自動計算。未登録だと請求が「要入力」になります。")
cli_cols = [
    ("取引先コード", 11, "in", "@", None, "重複不可。請求台帳ではこのコードで選択します"),
    ("取引先名", 40, "in", None, None, "支店・工事区分ごとに請求書を分けている場合は、それぞれ別の取引先として登録"),
    ("振込名義(カナ)", 22, "in", None, None, "通帳に表示される名義。入金明細での照合に使います"),
    ("締日(参考)", 9, "in", None, None, "例: 末日 / 20日"),
    ("支払サイト(月)", 10, "in", "0", None, "請求月の何か月後に支払われるか。当月=0、翌月=1、翌々月=2"),
    ("支払日", 8, "in", "0", None, "1〜31。末日払いは31(月の日数に合わせて自動調整)"),
    ("毎月請求", 8, "in", None, None, "毎月必ず請求する取引先は「有」。対象月の請求が未登録だと「請求漏れ候補」"),
    ("既定の送付方法", 16, "in", None, None, None),
    ("送付先(メール/住所/システム)", 28, "in", None, None, None),
    ("担当者", 10, "in", None, None, "督促・状況確認の担当者"),
    ("備考", 48, "in", None, None, None),
    ("重複チェック", 11, "auto", None,
     f'=IF({cell(C, "code")}="","",IF(COUNTIF({CLI["code"]},{cell(C, "code")})>1,"コード重複",""))', None),
]
table(cm, cli_cols, LAST["cli"])
add_list(cm, f"{C['monthly']}{FIRST}:{C['monthly']}{LAST['cli']}", list_src("I"))
add_list(cm, f"{C['method']}{FIRST}:{C['method']}{LAST['cli']}", list_src("F"))
validation(cm, f"{C['site']}{FIRST}:{C['site']}{LAST['cli']}",
           DataValidation(type="whole", operator="between", formula1="0", formula2="12", allow_blank=True),
           "0〜12の整数を入力してください")
validation(cm, f"{C['payday']}{FIRST}:{C['payday']}{LAST['cli']}",
           DataValidation(type="whole", operator="between", formula1="1", formula2="31", allow_blank=True),
           "1〜31の整数を入力してください(末日=31)")
highlight(cm, f"{C['dup']}{FIRST}:{C['dup']}{LAST['cli']}", f'${C["dup"]}{FIRST}<>""', "FFFF00", "9C0006")
highlight(cm, f"{C['site']}{FIRST}:{C['payday']}{LAST['cli']}",
          f'AND(${C["name"]}{FIRST}<>"",OR(${C["site"]}{FIRST}="",${C["payday"]}{FIRST}=""))', "F8CBAD", "9C0006")

clients = []  # (コード, 名称, カナ, 締日, サイト, 支払日, 毎月, 送付方法, 送付先, 担当, 備考)
if args.sample:
    clients = [
        ("C001", "株式会社サンプル商事", "カ)サンプルシヨウジ", "末日", 1, 31, "有", "メール(PDF)",
         "keiri@sample.example", "山田", "サンプル(運用前に削除)"),
        ("C002", "有限会社テスト工業", "ユ)テストコウギヨウ", "末日", 2, 10, "無", "郵送", "東京都○○区…", "佐藤",
         "サンプル(運用前に削除)"),
        ("C003", "合同会社デモ", "ド)デモ", "20日", 1, 31, "有", "請求書発行システム", "", "山田",
         "サンプル(運用前に削除)"),
    ]
if args.src:
    import os
    import sys
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    from current_sheet import customers, load, payment_terms, terms_text

    _, wb_v, src_sheets = load(args.src)
    terms = payment_terms(wb_v, src_sheets)
    latest = src_sheets[-1][0]
    prev = src_sheets[-2][0] if len(src_sheets) > 1 else latest
    for i, (name, info) in enumerate(customers(wb_v, src_sheets), start=1):
        est = terms.get(name)
        notes = []
        site = day = None
        if est:
            lag, d, cnt, conf = est
            notes.append(f"支払条件は入金実績{cnt}件からの推定({terms_text(lag, d)}・確からしさ{conf})→要確認")
            if conf in ("高", "中"):
                site, day = lag, d
        else:
            notes.append("入金実績なし→支払条件を入力してください")
        months = info["請求月数"].get(prev, 0)
        monthly = "有" if months >= 10 else None
        if monthly:
            notes.append(f"{prev}は{months}か月請求あり→毎月請求と推定")
        used = [s for s in info["期"] if info["請求月数"].get(s)]
        if not used:
            notes.append("現行表に請求実績なし")
        elif latest not in used and prev not in used:
            notes.append(f"最終請求: {used[-1]}")
        clients.append((f"T{i:03d}", name, None, None, site, day, monthly, None, None, None, "。".join(notes)))

for i, vals in enumerate(clients):
    for key, v in zip(CLI_KEYS, vals):
        if v is not None:
            cm[f"{C[key]}{FIRST + i}"] = v
protect(cm)

# ---------------------------------------------------------------- 請求台帳
iv = wb.create_sheet("請求台帳")
title(iv, "請求台帳",
      "請求書1枚=1行。請求書を作る前に登録し、採番された請求番号を請求書に記載。黄色=入力 / 灰色=自動。行は削除せず、誤りは「取消」を選択。")


def x(key):
    return cell(I, key)


def cli_lookup(key):
    return f'INDEX({CLI[key]},MATCH({x("code")},{CLI["code"]},0))'


SITE, PAYDAY = cli_lookup("site"), cli_lookup("payday")
due_auto = (f'=IF(OR({x("date")}="",{x("code")}="",{x("name")}="コード誤り"),"",IFERROR('
            f'IF(OR({SITE}="",{PAYDAY}=""),"",DATE(YEAR({x("date")}),MONTH({x("date")})+{SITE},'
            f'MIN({PAYDAY},DAY(EOMONTH({x("date")},{SITE}))))),""))')
missing = ('=IF({no}="","",TRIM('
           'IF({date}="","請求日 ","")'
           '&IF({code}="","取引先 ",IF({name}="コード誤り","取引先コード誤り ",""))'
           '&IF({subject}="","件名 ","")'
           '&IF({amount}="","請求額 ","")'
           '&IF(AND({amount}<>"",{due}=""),"入金期日(マスタの支払条件が未登録) ","")'
           '&IF({method}="","送付方法 ","")'
           '&IF(AND({sent}<>"",{sender}=""),"送付者 ","")'
           '&IF(AND({sent}<>"",N({sent})>{base}),"送付日が未来の日付 ","")'
           '&IF(AND({sent}<>"",N({paid})>0,N({bal})<=0,{checker}=""),"消込確認者 ","")))').format(
    base=BASE_DATE, **{k: x(k) for k in I})
status = ('=IF({no}="","",IF({cancel}="取消","取消",IF({missing}<>"","要入力",IF({sent}="","未送付",'
          'IF({bal}=0,"入金済",IF({bal}<0,"過入金",'
          'IF({base}>{due},IF(N({paid})>0,"一部入金・超過","期日超過"),'
          'IF(N({paid})>0,"一部入金",IF({due}-{base}<={near},"期日間近","入金待ち")))))))))').format(
    base=BASE_DATE, near=NEAR_DAYS, **{k: x(k) for k in I})
inv_cols = [
    ("請求番号", 12, "auto", None,
     f'=IF(COUNTA({x("date")},{x("code")},{x("subject")},{x("amount")})=0,"",'
     f'{PREFIX}&TEXT(ROW()-{FIRST - 1},"00000"))',
     "請求日・取引先・件名・請求額のいずれかを入力すると自動採番。この番号を請求書に記載します"),
    ("請求日", 11, "in", FMT_DATE, None, "請求書の日付(締日)。入金期日の起算月になります。年から入力"),
    ("取引先コード", 10, "in", "@", None, "取引先マスタのコードから選択"),
    ("取引先名", 30, "auto", None,
     f'=IF({x("code")}="","",IFERROR({cli_lookup("name")},"コード誤り"))', None),
    ("件名", 28, "in", None, None, "例: 2026年10月分 検査費 / ○○邸 外部付帯工事"),
    ("請求額(税込)", 13, "in", FMT_YEN, None, "請求書1枚の税込金額。複数の請求書を合算しない"),
    ("入金期日(自動)", 12, "auto", FMT_DATE, due_auto, "請求日と取引先マスタの支払条件から自動計算"),
    ("入金期日(個別)", 12, "in", FMT_DATE, None, "個別に期日が決まっている場合のみ入力(自動より優先)"),
    ("入金期日", 12, "auto", FMT_DATE, f'=IF({x("due_manual")}<>"",{x("due_manual")},{x("due_auto")})',
     "この期日でステータスを判定"),
    ("送付方法", 14, "in", None, None, None),
    ("送付日", 11, "in", FMT_DATE, None, "送付(先方システム登録)したら必ず入力。空欄の間は「未送付」"),
    ("送付者", 10, "in", None, None, None),
    ("入金額", 12, "auto", FMT_YEN, f'=IF({x("no")}="","",SUMIF({DEP["inv"]},{x("no")},{DEP["amount"]}))',
     "入金明細で、この請求番号に消込した金額の合計"),
    ("最終入金日", 11, "auto", FMT_DATE,
     f'=IF(OR({x("no")}="",N({x("paid")})=0),"",SUMPRODUCT(MAX(({DEP["inv"]}={x("no")})*{DEP["date"]})))', None),
    ("手数料・値引等", 11, "in", FMT_YEN, None, "先方が振込手数料を差し引いた場合などの差額。入れると残高から控除"),
    ("未回収残高", 12, "auto", FMT_YEN,
     f'=IF(OR({x("no")}="",{x("amount")}=""),"",{x("amount")}-N({x("paid")})-N({x("adj")}))', None),
    ("取消", 7, "in", None, None, "誤発行・再発行などで無効にする請求は「取消」を選び、備考に理由を記入(行は削除しない)"),
    ("ステータス", 13, "auto", None, status, "「はじめに」シートの4.を参照"),
    ("超過日数", 8, "auto", "0",
     f'=IF(OR({x("status")}="期日超過",{x("status")}="一部入金・超過"),{BASE_DATE}-{x("due")},"")', None),
    ("入力漏れ", 26, "auto", None, missing, "未入力・誤りの項目名を表示。空欄になればOK"),
    ("消込確認者", 10, "in", None, None, "入金済になった請求を確認した人(入金明細の記帳者とは別の人が望ましい)"),
    ("備考", 30, "in", None, None, None),
]
table(iv, inv_cols, LAST["inv"])
rng_i = lambda k: f"{I[k]}{FIRST}:{I[k]}{LAST['inv']}"  # noqa: E731
for k in ("date", "due_manual", "sent"):
    add_date(iv, rng_i(k))
for k in ("amount", "adj"):
    add_number(iv, rng_i(k))
add_list(iv, rng_i("code"), CLI["code"], "取引先マスタに登録されたコードを選択してください")
add_list(iv, rng_i("method"), list_src("F"))
add_list(iv, rng_i("cancel"), list_src("H"))
status_colors(iv, I["status"], LAST["inv"])
highlight(iv, rng_i("missing"), f'${I["missing"]}{FIRST}<>""', "FFFF00", "9C0006")
iv.conditional_formatting.add(f"A{FIRST}:{I['cancel']}{LAST['inv']}", FormulaRule(
    formula=[f'${I["cancel"]}{FIRST}="取消"'], font=Font(name=FONT, strike=True, color="7F7F7F")))

if args.sample:
    # 請求日, コード, 件名, 請求額, 個別期日, 送付方法, 送付日, 送付者, 手数料, 取消, 確認者, 備考
    sample_inv = [
        (date(2026, 7, 31), "C001", "2026年7月分 保守料", 110000, None, "メール(PDF)", date(2026, 7, 31), "鈴木",
         None, None, "田中", "SAMPLE: 入金済"),
        (date(2026, 7, 31), "C002", "○○部品 納品分", 275000, None, "郵送", date(2026, 8, 3), "鈴木", 440, None,
         None, "SAMPLE: 手数料差引で入金(確認者未記入→要入力)"),
        (date(2026, 6, 30), "C002", "○○部品 6月納品分", 165000, None, "郵送", date(2026, 7, 1), "鈴木", None, None,
         None, "SAMPLE: 期日超過"),
        (date(2026, 8, 31), "C003", "8月分 業務委託料", 88000, None, "請求書発行システム", date(2026, 9, 1), "鈴木",
         None, None, None, "SAMPLE: 一部入金"),
        (date(2026, 9, 30), "C001", "2026年9月分 保守料", 110000, None, "メール(PDF)", None, None, None, None,
         None, "SAMPLE: 未送付"),
        (date(2026, 9, 15), "C003", "追加作業費", 33000, None, "請求書発行システム", date(2026, 9, 15), "鈴木", None,
         "取消", None, "SAMPLE: 金額誤りのため取消→次行で再発行"),
        (date(2026, 9, 15), "C003", "追加作業費(再発行)", 36300, date(2026, 10, 5), "請求書発行システム",
         date(2026, 9, 16), "鈴木", None, None, None, "SAMPLE: 個別期日・期日間近"),
    ]
    keys = ["date", "code", "subject", "amount", "due_manual", "method", "sent", "sender", "adj", "cancel",
            "checker", "note"]
    for i, vals in enumerate(sample_inv):
        for k, v in zip(keys, vals):
            if v is not None:
                iv[f"{I[k]}{FIRST + i}"] = v
protect(iv)

# ---------------------------------------------------------------- 入金明細
dp = wb.create_sheet("入金明細")
title(dp, "入金明細",
      "通帳・ネットバンキングの入金を全件、1入金1行で転記し「消込先請求番号」を入力。1つの入金で複数の請求分が払われた場合は行を分けて金額を按分。")


def y(key):
    return cell(D, key)


def inv_lookup(key):
    return f'IFERROR(INDEX({INV[key]},MATCH({y("inv")},{INV["no"]},0)),"")'


used = f'COUNTA({y("date")},{y("payer")},{y("amount")},{y("inv")})=0'
dep_state = (f'=IF({used},"",IF(OR({y("date")}="",{y("amount")}=""),"要入力",'
             f'IF(N({y("date")})>{BASE_DATE},"日付誤り",IF({y("inv")}="","未消込",'
             f'IF(ISNA(MATCH({y("inv")},{INV["no"]},0)),"番号誤り",'
             f'IF(INDEX({INV["cancel"]},MATCH({y("inv")},{INV["no"]},0))="取消","取消請求へ消込","消込済"))))))')
dep_cols = [
    ("No", 6, "auto", None, f'=IF({used},"",ROW()-{FIRST - 1})', None),
    ("入金日", 11, "in", FMT_DATE, None, "通帳の日付。年から入力(未来の日付は「日付誤り」)"),
    ("入金口座", 16, "in", None, None, None),
    ("振込名義", 24, "in", None, None, "通帳の表示どおりに入力"),
    ("入金額", 12, "in", FMT_YEN, None, "通帳の実額(請求額を写さない)"),
    ("消込先請求番号", 13, "in", None, None, "対応する請求台帳の請求番号。不明なら空欄のまま(=未消込として残る)"),
    ("取引先コード", 10, "auto", None, f'=IF({y("inv")}="","",{inv_lookup("code")})', None),
    ("取引先名(請求)", 30, "auto", None, f'=IF({y("inv")}="","",{inv_lookup("name")})',
     "消込先請求の取引先。振込名義と一致しているか目視で確認"),
    ("請求額", 12, "auto", FMT_YEN, f'=IF({y("inv")}="","",{inv_lookup("amount")})', None),
    ("消込状態", 13, "auto", None, dep_state, None),
    ("記帳者", 10, "in", None, None, None),
    ("備考", 30, "in", None, None, None),
]
table(dp, dep_cols, LAST["dep"])
rng_d = lambda k: f"{D[k]}{FIRST}:{D[k]}{LAST['dep']}"  # noqa: E731
add_date(dp, rng_d("date"))
add_number(dp, rng_d("amount"))
add_list(dp, rng_d("account"), list_src("G"))
status_colors(dp, D["state"], LAST["dep"])
if args.sample:
    sample_dep = [
        (date(2026, 8, 31), "平塚信用金庫", "カ)サンプルシヨウジ", 110000, "INV-00001", "佐々木", "SAMPLE"),
        (date(2026, 9, 10), "平塚信用金庫", "ユ)テストコウギヨウ", 274560, "INV-00002", "佐々木",
         "SAMPLE: 振込手数料440円差引"),
        (date(2026, 9, 28), "平塚信用金庫", "ド)デモ", 50000, "INV-00004", "佐々木", "SAMPLE: 一部入金"),
        (date(2026, 9, 28), "平塚信用金庫", "ヤマダ タロウ", 5500, None, "佐々木", "SAMPLE: 不明入金(未消込)"),
    ]
    for i, vals in enumerate(sample_dep):
        for k, v in zip(["date", "account", "payer", "amount", "inv", "clerk", "note"], vals):
            if v is not None:
                dp[f"{D[k]}{FIRST + i}"] = v
protect(dp)

# ---------------------------------------------------------------- チェック
ck = wb.create_sheet("チェック")
title(ck, "チェック(ダッシュボード)",
      "判定が「要対応」の項目を解消してください。月次締めでは全項目「OK」を確認して「月次チェック記録」に記入します。")
for col, w in zip("ABCDEF", (3, 40, 12, 16, 11, 64)):
    ck.column_dimensions[col].width = w
ck["B4"], ck["B4"].font = "対象月(月次チェック用)", F_BOLD
ck["C4"] = f"=DATE(YEAR({BASE_DATE}),MONTH({BASE_DATE})-1,1)"
ck["C4"].number_format = FMT_MONTH
ck["C4"].fill, ck["C4"].border, ck["C4"].font, ck["C4"].protection = FILL_INPUT, BORDER, F_BOLD, UNLOCKED
ck["D4"], ck["D4"].font = "既定は基準日の前月。確認したい月の1日を入力(例: 2026/10/1)", F_NOTE
ck["B5"], ck["B5"].font = "基準日", F_BOLD
ck["C5"] = f"={BASE_DATE}"
ck["C5"].number_format, ck["C5"].font = FMT_DATE, F_BOLD
ck["D5"], ck["D5"].font = "「設定」シートで変更", F_NOTE
for i, h in enumerate(["チェック項目", "件数", "金額", "判定", "対応方法"], start=2):
    c = ck.cell(row=7, column=i, value=h)
    c.font, c.fill, c.alignment, c.border = F_HEAD, FILL_HEAD, CENTER, BORDER


def cnt(s):
    return f'COUNTIF({INV["status"]},"{s}")'


def amt(s):
    return f'SUMIF({INV["status"]},"{s}",{INV["bal"]})'


def dcnt(s):
    return f'COUNTIF({DEP["state"]},"{s}")'


def damt(s):
    return f'SUMIF({DEP["state"]},"{s}",{DEP["amount"]})'


month_end = f"EOMONTH({TARGET_MONTH},0)"
checks = [
    ("section", "■ 請求(請求台帳)"),
    ("要入力(入力漏れ)", "=" + cnt("要入力"), None, True,
     "請求台帳を「ステータス=要入力」で絞り込み、「入力漏れ」列の項目を入力"),
    ("未送付", "=" + cnt("未送付"), f'=SUMIF({INV["status"]},"未送付",{INV["amount"]})', True,
     "請求書を送付し送付日・送付者を入力。不要な請求なら「取消」"),
    ("期日超過(一部入金含む)", f"={cnt('期日超過')}+{cnt('一部入金・超過')}",
     f"={amt('期日超過')}+{amt('一部入金・超過')}", True,
     "担当者から先方へ確認・督促。入金済なら入金明細の転記・消込漏れを確認"),
    ("過入金", "=" + cnt("過入金"), "=" + amt("過入金"), True,
     "二重振込・消込先の誤りを確認。返金または次回請求と相殺し備考に記録"),
    ("期日間近(参考)", "=" + cnt("期日間近"), "=" + amt("期日間近"), False, "期日前の確認・リマインド(任意)"),
    ("入金待ち(参考)", f"={cnt('入金待ち')}+{cnt('一部入金')}", f"={amt('入金待ち')}+{amt('一部入金')}", False,
     "対応不要"),
    ("section", "■ 入金(入金明細)"),
    ("未消込の入金", "=" + dcnt("未消込"), "=" + damt("未消込"), True,
     "振込名義・金額から該当請求を探して請求番号を入力。不明なら先方・担当者に確認"),
    ("消込先の請求番号誤り", "=" + dcnt("番号誤り"), "=" + damt("番号誤り"), True, "請求番号を正しく入力し直す"),
    ("取消した請求への消込", "=" + dcnt("取消請求へ消込"), "=" + damt("取消請求へ消込"), True,
     "再発行後の請求番号へ付け替える"),
    ("入金日の誤り(未来の日付)", "=" + dcnt("日付誤り"), "=" + damt("日付誤り"), True,
     "入金日を年から入力し直す(「12/30」だけの入力は年がずれる)"),
    ("入金明細の入力漏れ", "=" + dcnt("要入力"), None, True, "入金日・入金額を入力"),
    ("section", "■ 請求漏れ(対象月)"),
    ("毎月請求先で対象月の請求が未登録", f'=COUNTIF({SHEET["bal"]}!$J${FIRST}:$J${LAST["cli"]},"請求漏れ候補")',
     None, True, "「取引先別残高」の請求漏れ候補を確認し、請求するか理由を記録"),
    ("マスタの取引先コード重複", f'=COUNTIF({CLI["dup"]},"コード重複")', None, True, "取引先マスタを修正"),
    ("section", "■ 月次の数字(対象月)"),
    ("対象月の請求(取消除く)",
     f'=COUNTIFS({INV["date"]},">="&{TARGET_MONTH},{INV["date"]},"<="&{month_end},{INV["cancel"]},"<>取消")',
     f'=SUMIFS({INV["amount"]},{INV["date"]},">="&{TARGET_MONTH},{INV["date"]},"<="&{month_end},'
     f'{INV["cancel"]},"<>取消")', None, "工事・検査の完了一覧、作業報告と件数を突合し、請求漏れがないか確認"),
    ("対象月の入金",
     f'=COUNTIFS({DEP["date"]},">="&{TARGET_MONTH},{DEP["date"]},"<="&{month_end})',
     f'=SUMIFS({DEP["amount"]},{DEP["date"]},">="&{TARGET_MONTH},{DEP["date"]},"<="&{month_end})', None,
     "通帳の対象月の入金合計(売上入金分)と一致するか確認"),
    ("未回収残高 合計(基準日時点)",
     f'=COUNTIFS({INV["bal"]},">0",{INV["cancel"]},"<>取消")+COUNTIFS({INV["bal"]},"<0",{INV["cancel"]},"<>取消")',
     f'=SUMIF({INV["cancel"]},"<>取消",{INV["bal"]})', None, "下の会計ソフトの売掛金残高と照合"),
]
r = 8
row_of = {}
for item in checks:
    if item[0] == "section":
        ck.cell(row=r, column=2, value=item[1]).font = F_BOLD
        for col in range(2, 7):
            ck.cell(row=r, column=col).fill = FILL_SECTION
        r += 1
        continue
    label, count_f, amount_f, judged, action = item
    row_of[label] = r
    ck.cell(row=r, column=2, value=label)
    ck.cell(row=r, column=3, value=count_f).number_format = '#,##0"件"'
    if amount_f:
        ck.cell(row=r, column=4, value=amount_f).number_format = FMT_YEN
    if judged is True:
        ck.cell(row=r, column=5, value=f'=IF(C{r}=0,"OK","要対応")')
    elif judged is False:
        ck.cell(row=r, column=5, value="参考")
    ck.cell(row=r, column=6, value=action)
    for col in range(2, 7):
        c = ck.cell(row=r, column=col)
        c.font, c.border = F_BASE, BORDER
        if col == 6:
            c.alignment = WRAP
    r += 1
r += 1
ck.cell(row=r, column=2, value="■ 会計ソフトとの残高照合(月次)").font = F_BOLD
for col in range(2, 7):
    ck.cell(row=r, column=col).fill = FILL_SECTION
r += 1
acct_row = r
ck.cell(row=r, column=2, value="会計ソフトの売掛金残高(基準日時点)")
acct = ck.cell(row=r, column=4)
acct.fill, acct.number_format, acct.protection = FILL_INPUT, FMT_YEN, UNLOCKED
ck.cell(row=r, column=6, value="会計ソフト(または税理士への資料)の売掛金残高を入力(基準日と同じ日付時点)")
r += 1
ck.cell(row=r, column=2, value="差額(台帳 − 会計ソフト)")
ck.cell(row=r, column=4, value=f'=IF(D{acct_row}="","",D{row_of["未回収残高 合計(基準日時点)"]}-D{acct_row})'
        ).number_format = FMT_YEN
ck.cell(row=r, column=5, value=f'=IF(D{acct_row}="","未入力",IF(D{r}=0,"OK","要対応"))')
ck.cell(row=r, column=6, value="差額があれば、台帳への登録漏れ・会計側の計上漏れ・消込漏れのいずれかを調査")
for rr in (acct_row, r):
    for col in range(2, 7):
        c = ck.cell(row=rr, column=col)
        c.font, c.border = F_BASE, BORDER
highlight(ck, f"E8:E{r}", '$E8="要対応"', "F8CBAD", "9C0006")
highlight(ck, f"E8:E{r}", '$E8="OK"', "E2EFDA", "375623")
highlight(ck, f"E8:E{r}", '$E8="未入力"', "FFFF00", "9C0006")
r += 2
ck.cell(row=r, column=2, value="総合判定").font = F_BOLD
tot = ck.cell(row=r, column=3,
              value=f'=IF(COUNTIF(E8:E{r - 2},"要対応")+COUNTIF(E8:E{r - 2},"未入力")=0,"すべてOK","要対応あり")')
tot.font, tot.border = F_BOLD, BORDER
highlight(ck, f"C{r}", f'$C${r}="要対応あり"', "F8CBAD", "9C0006")
highlight(ck, f"C{r}", f'$C${r}="すべてOK"', "E2EFDA", "375623")
protect(ck)

# ---------------------------------------------------------------- 取引先別残高
bl = wb.create_sheet("取引先別残高")
title(bl, "取引先別残高",
      "取引先マスタの順に自動表示(入力不要)。超過日数別の残高と、毎月請求先の対象月請求漏れ候補を表示します。")
M = SHEET["cli"]


def by_code(extra):
    return f'SUMIFS({INV["bal"]},{INV["code"]},A#,{extra})'


bal_cols = [
    ("取引先コード", 11, "auto", None, f'=IF({M}!{C["code"]}#="","",{M}!{C["code"]}#)', None),
    ("取引先名", 34, "auto", None, f'=IF(A#="","",{M}!{C["name"]}#)', None),
    ("担当者", 10, "auto", None, f'=IF(A#="","",{M}!{C["rep"]}#&"")', None),
    ("未回収残高", 13, "auto", FMT_YEN, '=IF(A#="","",' + by_code(INV["cancel"] + ',"<>取消"') + ")", None),
    ("期日内", 12, "auto", FMT_YEN, '=IF(A#="","",D#-SUM(F#:H#))', None),
    ("超過 1〜30日", 12, "auto", FMT_YEN,
     '=IF(A#="","",' + by_code(INV["over"] + ',">=1",' + INV["over"] + ',"<=30"') + ")", None),
    ("超過 31〜60日", 12, "auto", FMT_YEN,
     '=IF(A#="","",' + by_code(INV["over"] + ',">=31",' + INV["over"] + ',"<=60"') + ")", None),
    ("超過 61日以上", 12, "auto", FMT_YEN, '=IF(A#="","",' + by_code(INV["over"] + ',">=61"') + ")", None),
    ("対象月の請求件数", 11, "auto", '0"件"',
     f'=IF(A#="","",COUNTIFS({INV["code"]},A#,{INV["date"]},">="&{TARGET_MONTH},'
     f'{INV["date"]},"<="&{month_end},{INV["cancel"]},"<>取消"))',
     "「チェック」シートの対象月に請求日がある請求の件数(取消除く)"),
    ("請求漏れ判定", 13, "auto", None, f'=IF(A#="","",IF(AND({M}!{C["monthly"]}#="有",I#=0),"請求漏れ候補",""))',
     "取引先マスタで毎月請求=有 なのに対象月の請求がない"),
]
table(bl, bal_cols, LAST["cli"])
bl.auto_filter.ref = None
tot_row = LAST["cli"] + 1
bl.cell(row=tot_row, column=2, value="合計").font = F_BOLD
bl["C3"], bl["C3"].font = "合計", F_BOLD
for col in "DEFGH":
    for rr, v in ((tot_row, f"=SUM({col}{FIRST}:{col}{LAST['cli']})"), (3, f"={col}{tot_row}")):
        c = bl[f"{col}{rr}"]
        c.value, c.font, c.number_format, c.border = v, F_BOLD, FMT_YEN, BORDER
highlight(bl, f"J{FIRST}:J{LAST['cli']}", f'$J{FIRST}="請求漏れ候補"', "F8CBAD", "9C0006")
for col in "FGH":
    bl.conditional_formatting.add(f"{col}{FIRST}:{col}{LAST['cli']}", FormulaRule(
        formula=[f"N(${col}{FIRST})>0"], font=Font(name=FONT, bold=True, color="C00000")))
protect(bl)

# ---------------------------------------------------------------- 月別集計
ms = wb.create_sheet("月別集計")
title(ms, "月別集計(取引先×月)",
      "請求額=請求日の月で集計(取消除く)、入金額=入金日の月で集計。収支一覧表の収入欄と同じ形です。期首は「設定」で変更。")
ms.column_dimensions["A"].width = 10
ms.column_dimensions["B"].width = 34
for i, h in enumerate(["取引先コード", "取引先名"], start=1):
    ms.merge_cells(start_row=3, start_column=i, end_row=HDR, end_column=i)
    c = ms.cell(row=3, column=i, value=h)
    c.font, c.fill, c.alignment, c.border = F_HEAD, FILL_HEAD_AUTO, CENTER, BORDER
LAST_MS = LAST["cli"]
unmatched_row, total_row = LAST_MS + 1, LAST_MS + 2
month_cols = []
for k in range(13):
    cb, cp = get_column_letter(3 + 2 * k), get_column_letter(4 + 2 * k)
    month_cols.append((cb, cp))
    ms.column_dimensions[cb].width = 12
    ms.column_dimensions[cp].width = 12
    ms.merge_cells(f"{cb}3:{cp}3")
    top = ms[f"{cb}3"]
    top.value = f"=DATE(YEAR({FY_CELL}),MONTH({FY_CELL})+{k},1)" if k < 12 else "年間合計"
    top.number_format = 'yyyy"年"m"月"'
    top.font, top.fill, top.alignment, top.border = F_HEAD, FILL_HEAD, CENTER, BORDER
    for col, label in ((cb, "請求額"), (cp, "入金額")):
        h = ms[f"{col}{HDR}"]
        h.value, h.font, h.fill, h.alignment, h.border = label, F_HEAD, FILL_HEAD, CENTER, BORDER
for r in range(FIRST, total_row + 1):
    is_client = r <= LAST_MS
    if is_client:
        ms[f"A{r}"] = f'=IF({M}!{C["code"]}{r}="","",{M}!{C["code"]}{r})'
        ms[f"B{r}"] = f'=IF(A{r}="","",{M}!{C["name"]}{r})'
    elif r == unmatched_row:
        ms[f"B{r}"] = "未消込・番号誤りの入金"
    else:
        ms[f"B{r}"] = "合計"
    for k, (cb, cp) in enumerate(month_cols):
        m0 = f"{month_cols[min(k, 11)][0]}$3"
        lo, hi = f'">="&{m0}', f'"<"&DATE(YEAR({m0}),MONTH({m0})+1,1)'
        if k == 12:
            fb = f"=SUM({','.join(c[0] + str(r) for c in month_cols[:12])})"
            fp = f"=SUM({','.join(c[1] + str(r) for c in month_cols[:12])})"
        elif is_client:
            fb = (f'=IF($A{r}="","",SUMIFS({INV["amount"]},{INV["code"]},$A{r},{INV["date"]},{lo},'
                  f'{INV["date"]},{hi},{INV["cancel"]},"<>取消"))')
            fp = (f'=IF($A{r}="","",SUMIFS({DEP["amount"]},{DEP["code"]},$A{r},{DEP["date"]},{lo},'
                  f'{DEP["date"]},{hi}))')
        elif r == unmatched_row:
            fb = None
            fp = (f'=SUMPRODUCT(({DEP["code"]}="")*({DEP["date"]}>={m0})*'
                  f'({DEP["date"]}<DATE(YEAR({m0}),MONTH({m0})+1,1)),{DEP["amount"]})')
        else:
            fb = f"=SUM({cb}{FIRST}:{cb}{unmatched_row})"
            fp = f"=SUM({cp}{FIRST}:{cp}{unmatched_row})"
        for col, f in ((cb, fb), (cp, fp)):
            c = ms[f"{col}{r}"]
            c.value, c.number_format, c.border = f, FMT_YEN, BORDER
            c.font = F_BOLD if (not is_client or k == 12) else F_BASE
            c.fill = FILL_AUTO if is_client else FILL_SECTION
    for col in "AB":
        c = ms[f"{col}{r}"]
        c.font, c.border = (F_BASE if is_client else F_BOLD), BORDER
        c.fill = FILL_AUTO if is_client else FILL_SECTION
ms.freeze_panes = "C5"
protect(ms)

# ---------------------------------------------------------------- 月次チェック記録
mr = wb.create_sheet("月次チェック記録")
title(mr, "月次チェック記録",
      "毎月、締め後(例: 翌月5営業日以内)に実施。各項目を確認したら「済」を選択し、実施者・承認者を記入します。")
items = [
    ("対象月", 11, FMT_MONTH, None),
    ("①請求漏れ突合\n(工事・検査の完了と台帳)", 17, None, "工事完了・検査実施・作業報告の一覧と、対象月の請求台帳を1件ずつ突合"),
    ("②未送付 0件", 10, None, None),
    ("③要入力 0件", 10, None, None),
    ("④未消込入金 0件\n(番号誤り含む)", 13, None, None),
    ("⑤期日超過の督促・\n状況確認", 14, None, "期日超過の全件について、担当者が先方に確認し備考に結果を記録"),
    ("⑥会計ソフト残高\nとの照合(差額0)", 14, None, None),
    ("⑦毎月請求先の\n請求漏れ候補 0件", 14, None, None),
    ("実施者", 10, None, None),
    ("実施日", 11, FMT_DATE, None),
    ("承認者(責任者)", 12, None, "実施者とは別の人が、チェックシートの画面を見て承認"),
    ("承認日", 11, FMT_DATE, None),
    ("期日超過の件数・金額\n/ 特記事項", 40, None, None),
]
N_MONTHS = 36
for i, (head, width, fmt, comment) in enumerate(items, start=1):
    col = get_column_letter(i)
    c = mr.cell(row=HDR, column=i, value=head)
    c.font, c.fill, c.alignment, c.border = F_HEAD, FILL_HEAD, CENTER, BORDER
    if comment:
        c.comment = Comment(comment, "管理表")
    mr.column_dimensions[col].width = width
    for rr in range(FIRST, FIRST + N_MONTHS):
        cl = mr.cell(row=rr, column=i)
        cl.font, cl.fill, cl.border, cl.protection = F_BASE, FILL_INPUT, BORDER, UNLOCKED
        if fmt:
            cl.number_format = fmt
        if 2 <= i <= 8:
            cl.alignment = CENTER
mr.row_dimensions[HDR].height = 46
mr.freeze_panes = "B5"
for k in range(N_MONTHS):
    yy, mm = divmod(FY_START.month - 1 + k, 12)
    mr.cell(row=FIRST + k, column=1, value=date(FY_START.year + yy, mm + 1, 1))
last_mr = FIRST + N_MONTHS - 1
add_list(mr, f"B{FIRST}:H{last_mr}", list_src("J"))
add_date(mr, f"J{FIRST}:J{last_mr}")
add_date(mr, f"L{FIRST}:L{last_mr}")
mr.conditional_formatting.add(f"A{FIRST}:M{last_mr}", FormulaRule(
    formula=[f'AND($A{FIRST}<>"",$A{FIRST}<{TARGET_MONTH}+1,$L{FIRST}="")'],
    fill=PatternFill("solid", fgColor="FCE4D6", bgColor="FCE4D6")))
mr["N4"], mr["N4"].font = "※ 対象月までの行で承認日が空欄の行はオレンジ表示", F_NOTE
protect(mr)

# ---------------------------------------------------------------- 仕上げ
order = ["はじめに", "チェック", "請求台帳", "入金明細", "取引先別残高", "月別集計", "取引先マスタ",
         "月次チェック記録", "設定"]
wb._sheets = [wb[n] for n in order]
tabs = {"はじめに": "7F7F7F", "チェック": "C00000", "請求台帳": "1F4E78", "入金明細": "1F4E78",
        "取引先別残高": "548235", "月別集計": "548235", "取引先マスタ": "BF8F00", "月次チェック記録": "C00000",
        "設定": "7F7F7F"}
for s in wb.worksheets:
    s.sheet_properties.tabColor = tabs[s.title]

wb.save(args.out)
print("saved", args.out, f"(取引先 {len(clients)} 件)")
