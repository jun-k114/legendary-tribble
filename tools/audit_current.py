"""現行の収支一覧表を点検し、記入漏れ・不整合の一覧を Excel で出力するスクリプト。

使い方:  python tools/audit_current.py 現行の収支一覧表.xlsx [出力パス] [基準日 YYYY-MM-DD]
レイアウトの前提は current_sheet.py を参照。
"""
import re
import sys
from collections import Counter
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from current_sheet import (FEE_TOLERANCE, blank, load, month_of, num, payment_terms, terms_text,
                           ymd)

SRC = sys.argv[1]
OUT = sys.argv[2] if len(sys.argv) > 2 else "現行表チェック結果.xlsx"
TODAY = date.fromisoformat(sys.argv[3]) if len(sys.argv) > 3 else date.today()
DEFAULT_TERMS = (1, 31)  # 実績がない得意先は翌月末日払いと仮定

wb_f, wb_v, sheets = load(SRC)
terms = payment_terms(wb_v, sheets)


def is_formula(x):
    return isinstance(x, str) and x.startswith("=")


def fmt_date(x):
    d = ymd(x)
    return d.strftime("%Y/%m/%d") if d else (x if x is not None else "")


def due_date(name, y, mon):
    lag, day = terms.get(name, DEFAULT_TERMS)[:2]
    yy, mm = divmod(y * 12 + mon - 1 + lag, 12)
    mm += 1
    last = (date(yy + (mm == 12), mm % 12 + 1, 1) - date(yy, mm, 1)).days
    return date(yy, mm, min(day, last))


findings = []
summary = []
unpaid = []

for title, fy, sections in sheets:
    wf, wv = wb_f[title], wb_v[title]
    last_month = max((m for sec in sections.values() for r in range(sec[0], sec[1] + 1)
                      for m in range(12)
                      if any(not blank(wv.cell(r, 4 + 3 * m + k).value) for k in range(3))), default=-1)

    def add(sec, r, name, m, amt, dt, third, kind, level, text):
        y, mon = month_of(fy, m) if m is not None else (None, None)
        findings.append(dict(期=title, 区分=sec, 行=r, 名称=name, 月=f"{y}年{mon}月" if y else "",
                             金額=amt, 日付=fmt_date(dt), 実入金=third, 種類=kind, 重要度=level, 内容=text))

    for sec, (r0, r1) in sections.items():
        for r in range(r0, r1 + 1):
            if sec == "収入":
                name = wv.cell(r, 2).value
            else:
                name = " / ".join(str(x) for x in (wv.cell(r, 2).value, wv.cell(r, 3).value) if x) or None
            vals = [[wv.cell(r, 4 + 3 * m + k).value for k in range(3)] for m in range(12)]
            forms = [[wf.cell(r, 4 + 3 * m + k).value for k in range(3)] for m in range(12)]
            if all(blank(v) for mv in vals for v in mv):
                continue
            if blank(name):
                add(sec, r, "(名称なし)", None, None, None, None, "名称なし", "高",
                    "金額が入っているのに得意先/支出先名が空欄")
                name = f"(行{r})"
            name = name.strip()
            if not is_formula(wf.cell(r, 40).value) or (sec == "収入" and not is_formula(wf.cell(r, 41).value)):
                add(sec, r, name, None, None, None, None, "合計式なし", "高",
                    "AN(合計)/AO(実入金額合計)列に合計式がなく、年間合計から漏れる")
            filled, amounts = [], []
            billed = paid_sum = 0
            for m in range(12):
                amt, dt, third = vals[m]
                third_f = forms[m][2]
                if num(amt):
                    billed += amt
                    amounts.append(amt)
                if sec == "収入" and num(third):
                    paid_sum += third
                if blank(amt) and blank(dt) and blank(third):
                    continue
                filled.append(m)
                y, mon = month_of(fy, m)
                d = ymd(dt)
                if not blank(dt) and d is None:
                    add(sec, r, name, m, amt, dt, third, "日付が文字", "中",
                        f"日付が文字「{dt}」。複数回の入金/支払は1マスにまとめず、日付として扱えるよう記録を分ける必要がある")
                if d is not None:
                    lag = (d.year * 12 + d.month) - (y * 12 + mon)
                    if lag < -1 or lag > 6 or d > TODAY:
                        hint = ("「月/日」だけ入力するとExcelが入力した日の年を補うため、年をまたぐと誤りやすい"
                                if abs(lag) >= 10 else "")
                        add(sec, r, name, m, amt, dt, third, "日付の年/月誤り疑い", "高",
                            f"{y}年{mon}月分に対して日付が{fmt_date(d)}"
                            + ("(未来の日付)" if d > TODAY else f"({lag:+d}か月)") + "。" + hint)
                if sec == "収入":
                    if num(amt) and blank(dt) and blank(third):
                        due = due_date(name, y, mon)
                        est = terms.get(name)
                        basis = (f"推定期日 {due:%Y/%m/%d}({terms_text(*est[:2])}・実績{est[2]}件から推定)"
                                 if est else f"推定期日 {due:%Y/%m/%d}(実績がないため翌月末と仮定)")
                        if due < TODAY:
                            add(sec, r, name, m, amt, dt, third, "未入金または記入漏れ", "高",
                                f"金額はあるが入金日・実入金額が空欄。{basis}を経過。未入金なら督促、入金済みなら記入漏れ")
                        else:
                            add(sec, r, name, m, amt, dt, third, "入金期日前", "低",
                                f"{basis}。期日後に入金を確認")
                        unpaid.append([title, r, name, f"{y}年{mon}月", amt, due,
                                       "期日経過" if due < TODAY else "期日前"])
                    elif num(amt) and not blank(dt) and blank(third):
                        add(sec, r, name, m, amt, dt, third, "実入金額の記入漏れ", "高",
                            "入金日はあるが実入金額が空欄")
                    elif num(amt) and blank(dt) and not blank(third):
                        add(sec, r, name, m, amt, dt, third, "入金日の記入漏れ", "中",
                            "実入金額はあるが入金日が空欄")
                    if blank(amt) and not blank(third):
                        add(sec, r, name, m, amt, dt, third, "金額(請求額)の記入漏れ", "高",
                            "入金はあるが請求金額が空欄。この請求が収入合計に入っていない")
                    if num(amt) and num(third) and amt and third:
                        diff = amt - third
                        if 9.9 <= amt / third <= 10.1 or 9.9 <= third / amt <= 10.1:
                            add(sec, r, name, m, amt, dt, third, "桁違い疑い", "高",
                                f"金額と実入金額がちょうど約10倍違う(差額 {diff:,.0f}円)。0の打ち間違いの可能性")
                        elif abs(diff) > FEE_TOLERANCE:
                            add(sec, r, name, m, amt, dt, third, "金額と入金額の不一致", "中",
                                f"差額 {diff:,.0f}円。別の月分の入金を同じ欄に書いている/一部入金の可能性。"
                                "どの請求への入金か表から特定できない")
                    if is_formula(third_f) and re.fullmatch(r"=\$?[A-Z]+\$?\d+", third_f.replace(" ", "")):
                        add(sec, r, name, m, amt, dt, third, "実入金額が「=金額」の式", "低",
                            f"実入金額が「{third_f}」で金額をそのまま表示。通帳で実額を確認した記録にならない")
                else:
                    if num(amt) and blank(dt):
                        add(sec, r, name, m, amt, dt, third, "支払日の記入漏れ",
                            "中" if m < last_month else "低",
                            "金額はあるが支払日が空欄。未払いなら支払予定の確認、支払済みなら記入漏れ")
                    if blank(amt) and not blank(dt):
                        add(sec, r, name, m, amt, dt, third, "支払金額の記入漏れ", "高", "支払日はあるが金額が空欄")
            # 毎月請求がある得意先で途中の月だけ空欄
            if sec == "収入" and len(filled) >= 6:
                gaps = [m for m in range(min(filled), max(filled)) if m not in filled]
                fixed = amounts and Counter(amounts).most_common(1)[0][1] >= 4
                if gaps and len(gaps) <= 2:
                    for m in gaps:
                        add(sec, r, name, m, None, None, None, "毎月の請求が空欄", "高" if fixed else "中",
                            ("毎月同額の請求があるのにこの月だけ空欄。請求漏れ・記入漏れの可能性が高い" if fixed else
                             "前後の月は請求があるのにこの月だけ空欄。請求がなかったのか確認"))
            if sec == "収入":
                summary.append([title, r, name, billed, paid_sum, billed - paid_sum, wv.cell(r, 42).value or ""])
    upd = ymd(wv["AN2"].value)
    if upd and last_month >= 0:
        y, mon = month_of(fy, last_month)
        if upd.year * 12 + upd.month < y * 12 + mon:
            findings.append(dict(期=title, 区分="全体", 行=2, 名称="更新日", 月="", 金額=None, 日付=fmt_date(upd),
                                 実入金=None, 種類="更新日が古い", 重要度="低",
                                 内容=f"{y}年{mon}月分まで記入があるのに更新日が{fmt_date(upd)}のまま"))

# ------------------------------------------------------------------ 出力
FONT = "Meiryo UI"
F = Font(name=FONT, size=10)
FB = Font(name=FONT, size=10, bold=True)
FH = Font(name=FONT, size=10, bold=True, color="FFFFFF")
FT = Font(name=FONT, size=14, bold=True)
FN = Font(name=FONT, size=9, color="595959")
HEAD = PatternFill("solid", fgColor="1F4E78")
INPUT = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BD = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
LEVEL_FILL = {"高": "F8CBAD", "中": "FFE699", "低": "EDEDED"}
ORDER = {"高": 0, "中": 1, "低": 2}
YEN = "#,##0;[Red]-#,##0"

out = Workbook()


def write_table(ws, top, headers, rows, widths, fmts=None, input_cols=()):
    for i, h in enumerate(headers, 1):
        c = ws.cell(top, i, h)
        c.font, c.fill, c.border = FH, HEAD, BD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    for j, row in enumerate(rows, top + 1):
        for i, v in enumerate(row, 1):
            c = ws.cell(j, i, v)
            c.font, c.border = F, BD
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if fmts and fmts.get(i):
                c.number_format = fmts[i]
            if i in input_cols:
                c.fill = INPUT
    ws.freeze_panes = ws.cell(top + 1, 1)
    if rows:
        ws.auto_filter.ref = f"A{top}:{get_column_letter(len(headers))}{top + len(rows)}"


findings.sort(key=lambda f: (ORDER[f["重要度"]], f["種類"], f["期"], f["区分"], f["行"]))
kinds = Counter((f["種類"], f["重要度"]) for f in findings)
by_sheet = Counter((f["種類"], f["期"]) for f in findings)
titles = [s[0] for s in sheets]

ws = out.active
ws.title = "概要"
ws.sheet_view.showGridLines = False
ws["A1"] = "現行 収支一覧表 の点検結果"
ws["A1"].font = FT
ws["A2"] = (f"基準日 {TODAY:%Y/%m/%d}。自動点検のため誤検知も含みます。"
            "請求書控え・通帳と照らして「指摘一覧」の確認結果欄に記入してください。")
ws["A2"].font = FN
rows = [[lv, kind, n] + [by_sheet.get((kind, t), 0) for t in titles]
        for (kind, lv), n in sorted(kinds.items(), key=lambda x: (ORDER[x[0][1]], -x[1]))]
write_table(ws, 4, ["重要度", "指摘の種類", "件数"] + titles, rows, [8, 30, 8] + [8] * len(titles))
for j, row in enumerate(rows, 5):
    ws.cell(j, 1).fill = PatternFill("solid", fgColor=LEVEL_FILL[row[0]])
ws.freeze_panes = None
r = 6 + len(rows)
for t in [
    "■ 現行表の構造上の課題(漏れが起きる理由)",
    "・得意先×月の1マスに複数の請求を合算(例: =302500+224880)しているため、請求書1枚ごとの入金確認ができない",
    "・入金を別の月の欄や備考(例:「1月:11/14＝248,319」)に書いているため、どの請求が未入金か表から判別できない",
    "・実入金額が「=金額」の式の欄は、実際の入金額を確かめた記録にならず、差額や未入金が見えなくなる",
    "・日付を「月/日」だけで入力すると年が自動で補われ、年末年始をまたぐと翌年の日付になる(12/30 → 2026/12/30 等)",
    "・入金日を文字(「12/19、12/26」)で書いた欄は、日付として集計・期日管理ができない",
    "・入金期日・送付日の列がないため、「期日を過ぎた未入金」「作ったが送っていない請求書」を抽出できない",
    "",
    "■ 進め方",
    "1. 「未入金の可能性」シートの「期日経過」を最優先で、通帳・先方と確認",
    "2. 「指摘一覧」の重要度 高 → 中 の順に確認し、現行表を修正(確認結果欄に記録)",
    "3. 未入金と確定した請求は、新しい請求入金管理表の請求台帳に登録(移行)",
]:
    ws.cell(r, 1, t).font = FB if t.startswith("■") else F
    r += 1

ws = out.create_sheet("未入金の可能性")
unpaid.sort(key=lambda x: (x[6] != "期日経過", x[5]))
write_table(ws, 1, ["期", "行", "得意先", "請求月", "金額", "推定期日", "状態", "確認結果(入金済/未入金/請求なし)", "確認者"],
            unpaid, [6, 5, 40, 11, 13, 12, 10, 30, 10], fmts={5: YEN, 6: "yyyy/mm/dd"}, input_cols=(8, 9))
n = len(unpaid)
ws.cell(n + 3, 3, "※ 金額はあるが入金日・実入金額が空欄の欄。推定期日は過去の入金実績から推定した支払条件で計算。").font = FN
ws.cell(n + 4, 3, "※ 別の月の欄や備考に入金が書かれている場合もあるため、通帳で確認してください。").font = FN

ws = out.create_sheet("指摘一覧")
cols = ["期", "区分", "行", "名称", "月", "金額", "日付", "実入金", "種類", "重要度", "内容"]
rows = [[f[k] for k in cols] + ["", ""] for f in findings]
write_table(ws, 1, cols + ["確認結果", "確認者"], rows, [6, 6, 5, 34, 11, 12, 16, 12, 22, 7, 60, 30, 10],
            fmts={6: YEN, 8: YEN}, input_cols=(12, 13))
for j, row in enumerate(rows, 2):
    ws.cell(j, 10).fill = PatternFill("solid", fgColor=LEVEL_FILL[row[9]])

ws = out.create_sheet("得意先別の請求と入金")
write_table(ws, 1, ["期", "行", "得意先", "請求合計", "実入金合計", "差額(請求−入金)", "現行表の備考"], summary,
            [6, 5, 40, 14, 14, 16, 70], fmts={4: YEN, 5: YEN, 6: YEN})
n = len(summary)
ws.cell(n + 3, 3, "※ 差額がプラス=未入金または入金の記入漏れ、マイナス=前期分の入金・請求額の記入漏れ・桁違いの可能性。").font = FN
ws.cell(n + 4, 3, "※ 期をまたぐ入金(9月分が翌期に入金 等)があるため、前後の期と合わせて確認してください。").font = FN

ws = out.create_sheet("支払条件の推定")
rows = [[nm, lag, day, terms_text(lag, day), cnt, conf] for nm, (lag, day, cnt, conf) in sorted(terms.items())]
write_table(ws, 1, ["得意先", "支払サイト(月)", "支払日(31=末日)", "推定条件", "根拠件数", "確からしさ"], rows,
            [40, 12, 12, 14, 9, 10])
ws.cell(len(rows) + 3, 1, "※ 金額どおり入金された実績から、請求月→入金月のずれと入金日の最頻値で推定。25日以降は末日扱い。").font = FN

out.save(OUT)
print(f"saved {OUT}: {len(findings)} findings, unpaid candidates {len(unpaid)}")
for (kind, lv), n in sorted(kinds.items(), key=lambda x: (ORDER[x[0][1]], -x[1])):
    print(f"  [{lv}] {kind}: {n}")
